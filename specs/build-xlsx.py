# -*- coding: utf-8 -*-
import importlib.util, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

def load(p):
    sp=importlib.util.spec_from_file_location('x',p); mod=importlib.util.module_from_spec(sp)
    o=sys.stdout; sys.stdout=open(os.devnull,'w'); sp.loader.exec_module(mod); sys.stdout=o
    return mod.A, mod.M
A, OLD = load('map.py'); _, NEW = load('map_v2.py')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
_o=sys.stdout; sys.stdout=open(os.devnull,'w')
import review as RV
sys.stdout=_o

SITE='https://1-plus-1.ru'
GROUPS=[('Бухгалтерские услуги',['abon','vosst','decl','cons','malyy','sred','krup','glavbuh']),
        ('Финансовый консалтинг',['upruchet','audit','optnalog','sudexp']),
        ('Юридические услуги',['yurobsl','dolgi','pravaudit','egrul','reg','dogovory','ved','trud','sudy','intell','otnalog','yurcons']),
        ('Декларации 3-НДФЛ',['onlayn','ofis','ndfl3']),
        ('Отдельные услуги',['tz','voin'])]
ORDER=[c for _,cs in GROUPS for c in cs]
GRP={c:g for g,cs in GROUPS for c in cs}

DARK,OLIVE,OLIVEBG,ROSE,ROSEBG,GREY,LINE='1B1B1B','8E9142','EEF0DD','C32F55','FBE8EC','6C6D60','DCDCD2'
AMBER='B8860B'; AMBERBG='FDF4E3'
thin=Side(style='thin',color=LINE); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
PTAG={1:'① абонентское бухобслуживание',2:'② фин. и налоговый консалтинг',3:'③ юридическое обслуживание',0:'—'}

def hdr(ws,row,names,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for i,h in enumerate(names,1):
        c=ws.cell(row=row,column=i,value=h)
        c.font=Font(size=10,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=DARK)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
    ws.row_dimensions[row].height=32

wb=Workbook()

# ══ 1. Куда встаёт блок (как оставила Ольга) ══
ws=wb.active; ws.title='Куда встаёт блок'; ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=120
def put(r,t,s=11,b=False,col=DARK,h=None):
    c=ws.cell(row=r,column=2,value=t); c.font=Font(size=s,bold=b,color=col)
    c.alignment=Alignment(wrap_text=True,vertical='top')
    if h: ws.row_dimensions[r].height=h
put(2,'Блок «Другие услуги» на страницах услуг',20,True,DARK,30)
put(4,'КУДА ИМЕННО ВСТАЁТ',11,True,OLIVE,20)
put(5,'Блок ставится непосредственно перед секцией manager-callback — это форма обратного звонка и подвал. '
      'Правило одно для всех 29 страниц, исключений нет.',11,True,DARK,36)
put(6,'Почему здесь: человек дочитал страницу, предложение он уже видел. Дальше у него два сценария — оставить заявку '
      'или посмотреть соседнюю услугу. Блок не перехватывает внимание до оффера и не отодвигает форму так далеко, '
      'чтобы её перестали замечать.',11,False,DARK,50)
def place(path,row,w):
    im=PILImage.open(path); x=XLImage(path); x.width,x.height=w,int(im.height*w/im.width)
    ws.add_image(x,f'B{row}'); return row+x.height//20+2
r=place('shots/context.png',8,940)
put(r,'Мобильная версия (390px)',11,True,DARK,20)
r=place('shots/mobile.png',r+2,300)

# ══ 2. Проверка по критериям ══
t=wb.create_sheet('Проверка по критериям'); t.sheet_view.showGridLines=False
def cap(ws_,row,txt,s=10,b=False,col=GREY,h=16):
    c=ws_.cell(row=row,column=1,value=txt); c.font=Font(size=s,bold=b,color=col)
    c.alignment=Alignment(vertical='top'); ws_.row_dimensions[row].height=h
cap(t,1,'Проверка карты по двум критериям',16,True,DARK,24)
cap(t,2,'Критерий 1 — релевантность: логично ли человеку на этой странице перейти по ссылке.')
cap(t,3,'Критерий 2 — приоритет продажи: ① абонентское бухобслуживание, ② финансовый и налоговый консалтинг, ③ юридическое обслуживание.')
cap(t,4,'Вердикт «заменить» означает: ссылка релевантна, но ведёт туда, что нам продавать неинтересно. В колонке «Чем заменить» — предложение.',10,False,ROSE)
HR=6
hdr(t,HR,['№','Страница','Текущая ссылка','Поз.','Релевантность','Почему не высокая','Приоритет продажи','Вердикт','Чем заменить','Комментарий'],
      [4.5,30,30,5,16,40,30,13,34,40])
row=HR+1; pg=0
for code in ORDER:
    pg+=1; shade=pg%2==0
    old=OLD[code]; new=NEW[code]
    removed=[x for x in old if x not in new]
    added  =[x for x in new if x not in old]
    pair   ={rm: added[k] for k,rm in enumerate(removed) if k < len(added)}
    for i,tgt in enumerate(old,1):
        rel,why = RV.rel(code,tgt); p = RV.prio(tgt)
        keep = tgt in new
        repl = '' if keep else (A[pair[tgt]][1] + f'  ({RV.PNAME[RV.prio(pair[tgt])]})' if tgt in pair else '— убрать')
        vals=[pg,A[code][1],A[tgt][1],i,rel,why,PTAG[p],'оставить' if keep else 'заменить',repl if not keep else '','']
        for j,v in enumerate(vals,1):
            c=t.cell(row=row,column=j,value=v); c.border=BOX; c.font=Font(size=10)
            c.alignment=Alignment(vertical='center',wrap_text=(j in (2,3,6,7,9,10)))
            if shade: c.fill=PatternFill('solid',fgColor='F7F7F2')
            if j in (1,4): c.alignment=Alignment(horizontal='center',vertical='center')
            if j==5 and rel!='высокая': c.font=Font(size=10,color=AMBER); c.fill=PatternFill('solid',fgColor=AMBERBG)
            if j==7 and p: c.font=Font(size=10,bold=True,color=OLIVE)
            if j==8 and not keep: c.font=Font(size=10,bold=True,color=ROSE); c.fill=PatternFill('solid',fgColor=ROSEBG)
            if j==2 and i==1: c.font=Font(size=10,bold=True)
        t.row_dimensions[row].height=30; row+=1
t.freeze_panes=f'A{HR+1}'; t.auto_filter.ref=f'A{HR}:J{row-1}'

# ══ 3. Новая карта ══
n=wb.create_sheet('Новая карта'); n.sheet_view.showGridLines=False
cap(n,1,'Карта перелинковки после проверки',16,True,DARK,24)
cap(n,2,'Блок ставится непосредственно перед секцией manager-callback. 29 страниц × 4 ссылки = 116 ссылок.')
cap(n,3,'Приоритетных ссылок стало 64 из 116 вместо 42. Релевантность не ухудшилась.')
cap(n,4,'Жёлтым помечены новые ссылки — те, которых не было в прежней версии карты. Их 23 на 15 страницах.',10,False,AMBER)
HR2=6
hdr(n,HR2,['№','Раздел','Страница','URL страницы','№','Анкор ссылки','URL ссылки','Приоритет продажи','Новая','Статус','Комментарий'],
      [4.5,21,30,50,5,30,54,30,8,13,40])
row=HR2+1; pg=0
for code in ORDER:
    pg+=1; shade=pg%2==0; pu,pn=A[code]
    for i,tgt in enumerate(NEW[code],1):
        tu,tn=A[tgt]; p=RV.prio(tgt); fresh = tgt not in OLD[code]
        vals=[pg,GRP[code],pn,pu,i,tn,tu,PTAG[p],'да' if fresh else '','','']
        for j,v in enumerate(vals,1):
            c=n.cell(row=row,column=j,value=v); c.border=BOX; c.font=Font(size=10)
            c.alignment=Alignment(vertical='center',wrap_text=(j in (3,6,8,11)))
            if shade: c.fill=PatternFill('solid',fgColor='F7F7F2')
            if j in (1,5,9): c.alignment=Alignment(horizontal='center',vertical='center')
            if j in (4,7): c.font=Font(size=9,color='7A7A70'); c.hyperlink=SITE+v
            if j==8 and p: c.font=Font(size=10,bold=True,color=OLIVE)
            if j==9 and fresh: c.font=Font(size=10,bold=True,color=AMBER); c.fill=PatternFill('solid',fgColor=AMBERBG)
            if j==3 and i==1: c.font=Font(size=10,bold=True)
        n.row_dimensions[row].height=28; row+=1
LAST=row-1
n.freeze_panes=f'A{HR2+1}'; n.auto_filter.ref=f'A{HR2}:K{LAST}'
dv=DataValidation(type='list',formula1='"не начато,в работе,готово,вопрос"',allow_blank=True)
n.add_data_validation(dv); dv.add(f'J{HR2+1}:J{LAST}')

# ══ 4. Сводка ══
s=wb.create_sheet('Сводка'); s.sheet_view.showGridLines=False
for col,w in zip('ABCDE',(44,18,18,14,52)): s.column_dimensions[col].width=w
cap(s,1,'Что изменила проверка',16,True,DARK,24)
hdr(s,3,['Куда ведут ссылки','Было','Стало','Изм.','Комментарий'],[44,18,18,14,52])
from collections import Counter
co=Counter(RV.prio(x) for v in OLD.values() for x in v); cn=Counter(RV.prio(x) for v in NEW.values() for x in v)
rows=[(PTAG[1],co[1],cn[1],'главная коммерческая страница сайта'),
      (PTAG[2],co[2],cn[2],'вернули ссылки с бухгалтерских страниц — см. вопрос ниже'),
      (PTAG[3],co[3],cn[3],'растёт за счёт юридического раздела, он самый большой по числу страниц'),
      ('— вне приоритетов',co[0],cn[0],'было 64% всех ссылок, стало 45%')]
r=4
for nm,a_,b_,cm in rows:
    for j,v in enumerate([nm,a_,b_,f'{b_-a_:+d}',cm],1):
        c=s.cell(row=r,column=j,value=v); c.border=BOX; c.font=Font(size=10)
        c.alignment=Alignment(vertical='center',wrap_text=(j in (1,5)))
        if j in (2,3,4): c.alignment=Alignment(horizontal='center',vertical='center')
        if j==4: c.font=Font(size=10,bold=True,color=OLIVE if b_>=a_ else ROSE)
    s.row_dimensions[r].height=26; r+=1
r+=1
cap(s,r,'Входящие ссылки на приоритетные страницы',12,True,DARK,22); r+=1
hdr(s,r,['Страница','Было','Стало','Изм.','Комментарий'],[44,18,18,14,52]); r+=1
inbo=Counter(x for v in OLD.values() for x in v); inbn=Counter(x for v in NEW.values() for x in v)
for code in ['abon','optnalog','audit','upruchet','sudexp','yurobsl']:
    for j,v in enumerate([A[code][1],inbo[code],inbn[code],f'{inbn[code]-inbo[code]:+d}',PTAG[RV.prio(code)]],1):
        c=s.cell(row=r,column=j,value=v); c.border=BOX; c.font=Font(size=10)
        if j in (2,3,4): c.alignment=Alignment(horizontal='center')
        if j==4: c.font=Font(size=10,bold=True,color=OLIVE if inbn[code]>=inbo[code] else ROSE)
    s.row_dimensions[r].height=22; r+=1
r+=1
c=s.cell(row=r,column=1,value='ВОПРОС, КОТОРЫЙ НУЖНО РЕШИТЬ'); c.font=Font(size=12,bold=True,color=ROSE); r+=1
for line in ['В прошлой версии мы убрали финансовый консалтинг со страниц бухгалтерских услуг — по вашей правке.',
             'Но консалтинг стоит вторым в приоритете продаж, а бухгалтерский раздел — самый посещаемый.',
             'Чтобы приоритет работал, в новой карте эти 11 ссылок вернулись.',
             'Если решение убрать консалтинг из бухгалтерии было принципиальным — скажите, откачу.',
             'Тогда приоритет 2 останется с 15 ссылками вместо 26, и почти все — изнутри самого раздела консалтинга.']:
    c=s.cell(row=r,column=1,value=line); c.font=Font(size=10); c.alignment=Alignment(wrap_text=True)
    s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5); s.row_dimensions[r].height=16; r+=1

out='/home/user/Site1-1/specs/perelinkovka-drugie-uslugi.xlsx'
wb.save(out); print('сохранено:',out,os.path.getsize(out)//1024,'KB')
