# -*- coding: utf-8 -*-
import importlib.util, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def load(p):
    sp=importlib.util.spec_from_file_location('x',p); m=importlib.util.module_from_spec(sp)
    o=sys.stdout; sys.stdout=open(os.devnull,'w'); sp.loader.exec_module(m); sys.stdout=o
    return m
mm=load('map_v2.py'); A,M=mm.A,mm.M
T=load('texts.py')
SITE='https://1-plus-1.ru'
ORDER=['abon','vosst','decl','cons','malyy','sred','krup','glavbuh',
       'upruchet','audit','optnalog','sudexp',
       'yurobsl','dolgi','pravaudit','egrul','reg','dogovory','ved','trud','sudy','intell','otnalog','yurcons',
       'onlayn','ofis','ndfl3','tz','voin']
ACCENT={'vosst','audit','otnalog','glavbuh','cons','reg'}   # страницы с акцентной карточкой
DARK,LINE,ROSE,ROSEBG='1B1B1B','DCDCD2','C32F55','FBE8EC'
thin=Side(style='thin',color=LINE); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)

def head(ws,row,names,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for i,h in enumerate(names,1):
        c=ws.cell(row=row,column=i,value=h)
        c.font=Font(size=10,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=DARK)
        c.alignment=Alignment(horizontal='left',vertical='center'); c.border=BOX
    ws.row_dimensions[row].height=26

wb=Workbook()

# ── Лист 1: тексты услуг (утверждено) ──
ws=wb.active; ws.title='Тексты услуг'; ws.sheet_view.showGridLines=False
head(ws,1,['Услуга','Страница услуги','Текст'],[34,56,46])
r=2
for k,code in enumerate(ORDER):
    u,n=A[code]
    for j,v in enumerate([n, SITE+u, T.BASE[code]],1):
        c=ws.cell(row=r,column=j,value=v); c.border=BOX
        c.alignment=Alignment(vertical='center',wrap_text=(j in (1,3)))
        c.font=Font(size=9,color='7A7A70') if j==2 else Font(size=10)
        if j==2: c.hyperlink=v
        if k%2==1: c.fill=PatternFill('solid',fgColor='F7F7F2')
    ws.row_dimensions[r].height=24; r+=1
ws.freeze_panes='A2'

# ── Лист 2: все 116 карточек ──
f=wb.create_sheet('Все карточки'); f.sheet_view.showGridLines=False
head(f,1,['Страница, где стоит блок','Карточка','Текст под названием','Акцентная','Свой текст'],[32,32,46,11,11])
r=2
for k,src in enumerate(ORDER):
    for i,tgt in enumerate(M[src]):
        txt=T.text(src,tgt); own=(src,tgt) in T.OVER
        acc = 'да' if (src in ACCENT and i==0) else ''
        for j,v in enumerate([A[src][1], A[tgt][1], txt, acc, 'да' if own else ''],1):
            c=f.cell(row=r,column=j,value=v); c.border=BOX; c.font=Font(size=10)
            c.alignment=Alignment(vertical='center',wrap_text=(j in (1,2,3)))
            if j in (4,5):
                c.alignment=Alignment(horizontal='center',vertical='center')
                if v: c.font=Font(size=10,bold=True,color=ROSE); c.fill=PatternFill('solid',fgColor=ROSEBG)
            elif k%2==1: c.fill=PatternFill('solid',fgColor='F7F7F2')
        f.row_dimensions[r].height=24; r+=1
f.freeze_panes='A2'; f.auto_filter.ref=f'A1:E{r-1}'

out='/home/user/Site1-1/specs/teksty-kartochek.xlsx'
wb.save(out); print('сохранено:', out, os.path.getsize(out)//1024,'KB')
print('лист «Тексты услуг»: 29 | лист «Все карточки»: 116 | акцентных:', len(ACCENT))
